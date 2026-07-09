#!/usr/bin/env python3
"""Generate client-facing validation report DOCX."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

_REPO_ROOT = Path(__file__).resolve().parent.parent
OUTPUT_PATH = _REPO_ROOT / "final" / "validation_report.docx"


def _set_cell_shading(cell, fill: str) -> None:
    shading = cell._element.get_or_add_tcPr()
    shd = shading.makeelement(
        qn("w:shd"),
        {qn("w:fill"): fill, qn("w:val"): "clear"},
    )
    shading.append(shd)


def _style_heading(paragraph, *, size: int, bold: bool = True, color: RGBColor | None = None) -> None:
    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run(paragraph.text)
    if not paragraph.runs:
        paragraph.clear()
        run = paragraph.add_run(paragraph.text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    if color:
        run.font.color.rgb = color


def add_title_page(doc: Document) -> None:
    for _ in range(4):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Data Curation Validation Report")
    run.bold = True
    run.font.size = Pt(28)
    run.font.name = "Calibri"
    run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Indic Language Text-to-Speech Dataset")
    run.font.size = Pt(14)
    run.font.name = "Calibri"
    run.font.color.rgb = RGBColor(0x44, 0x54, 0x6A)

    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(f"Report Date: {date.today():%d %B %Y}")
    run.font.size = Pt(11)
    run.font.name = "Calibri"
    run.italic = True

    doc.add_page_break()


def add_executive_summary(doc: Document) -> None:
    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(
        "This report documents the two-phase validation framework applied to curated "
        "Indic-language text corpora intended for text-to-speech (TTS) synthesis. "
        "The framework ensures that every retained passage meets standards for "
        "linguistic quality, content safety, privacy compliance, and TTS suitability."
    )
    doc.add_paragraph(
        "Phase 1 focuses on content quality assurance through automated text cleaning, "
        "large language model (LLM) evaluation, and paragraph length standardization. "
        "Phase 2 applies specialized TTS-readiness filters to remove non-standard words, "
        "numerical expressions, citations, and structural artifacts that are difficult "
        "for speech synthesis systems to render naturally."
    )

    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    headers = [("Phase", "Objective"), ("Phase 1", "Content quality, safety, and paragraph sizing (300–450 words)"), ("Phase 2", "TTS readiness — removal of numbers, citations, and non-prose content")]
    for i, (col1, col2) in enumerate(headers):
        row = table.rows[i]
        row.cells[0].text = col1
        row.cells[1].text = col2
        if i == 0:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True
                _set_cell_shading(cell, "1F3864")
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    doc.add_paragraph()


def add_section_overview(doc: Document) -> None:
    doc.add_heading("1. Validation Framework Overview", level=1)
    doc.add_paragraph(
        "The validation process operates in two sequential phases. Each phase applies "
        "a defined set of checks and transformations. Records that fail critical quality "
        "or safety thresholds are excluded from the final dataset. Records that pass "
        "all checks are standardized to a target length of 300–450 words per paragraph."
    )

    steps = [
        ("Phase 1 — Step A", "Automated text normalization and cleaning"),
        ("Phase 1 — Step B", "LLM-based quality, sensitivity, and domain assessment"),
        ("Phase 1 — Step C", "Paragraph length filtering, pruning, and splitting"),
        ("Phase 2 — Step A", "Structural content filtering (lists, tables, boilerplate)"),
        ("Phase 2 — Step B", "Non-standard word (NSW) removal and text preprocessing"),
        ("Phase 2 — Step C", "Final TTS paragraph sizing and quality routing"),
    ]
    for label, desc in steps:
        p = doc.add_paragraph(style="List Number")
        p.add_run(f"{label}: ").bold = True
        p.add_run(desc)


def add_table_from_rows(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    *,
    col_widths: list[float] | None = None,
) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    for i, text in enumerate(headers):
        hdr.cells[i].text = text
        _set_cell_shading(hdr.cells[i], "D6E4F0")
        for p in hdr.cells[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(10)

    for r_idx, row_data in enumerate(rows):
        row = table.rows[r_idx + 1]
        for c_idx, text in enumerate(row_data):
            row.cells[c_idx].text = text
            for p in row.cells[c_idx].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)

    if col_widths:
        for row in table.rows:
            for i, width in enumerate(col_widths):
                row.cells[i].width = Inches(width)

    doc.add_paragraph()


def add_phase1(doc: Document) -> None:
    doc.add_heading("2. Phase 1 — Content Quality Validation", level=1)
    doc.add_paragraph(
        "Phase 1 establishes baseline content quality through deterministic text "
        "processing followed by LLM-assisted evaluation. The objective is to retain "
        "only coherent, safe, and linguistically appropriate prose passages."
    )

    doc.add_heading("2.1 Text Normalization and Cleaning", level=2)
    doc.add_paragraph(
        "Every text passage undergoes automated normalization before quality assessment. "
        "The following cleaning operations are applied:"
    )
    cleaning_rows = [
        ("1", "Unicode normalization", "Corrects encoding errors and applies standard Unicode normalization (NFC)"),
        ("2", "Control character removal", "Strips invisible, control, and private-use Unicode characters"),
        ("3", "Link removal", "Removes HTTP/HTTPS and www-prefixed URLs from body text"),
        ("4", "Emoji removal", "Strips Unicode emoji characters"),
        ("5", "Decorative artifact removal", "Removes repeated dot, dash, and bracket sequences from scraped content"),
        ("6", "Navigation boilerplate removal", "Strips serial-part navigation text and similar scraping artifacts"),
        ("7", "Foreign script filtering", "Removes words written in scripts other than the target language script"),
        ("8", "Latin word removal", "Removes pure Latin-script words embedded in Indic text"),
        ("9", "Bracket and quote balancing", "Corrects unmatched parentheses, brackets, and quotation marks"),
        ("10", "Terminal punctuation correction", "Adds language-appropriate sentence terminators where missing (danda for Indic scripts, Arabic full stop for Urdu/Sindhi)"),
        ("11", "Empty content exclusion", "Excludes passages that become empty after cleaning"),
    ]
    add_table_from_rows(doc, ["#", "Parameter", "Description"], cleaning_rows, col_widths=[0.4, 1.8, 4.3])

    doc.add_heading("2.2 Privacy and Personal Data Protection", level=2)
    doc.add_paragraph(
        "Personally identifiable information (PII) is detected using pattern matching and "
        "replaced with a standardized redaction marker. The passage is retained after "
        "redaction. The following PII categories are handled:"
    )
    pii_rows = [
        ("Email addresses", "Detected and redacted"),
        ("Indian mobile numbers", "10-digit numbers with optional +91 prefix — detected and redacted"),
        ("Aadhaar numbers", "12-digit national ID patterns — detected and redacted"),
        ("PAN numbers", "Indian tax identification patterns — detected and redacted"),
    ]
    add_table_from_rows(doc, ["PII Category", "Treatment"], pii_rows, col_widths=[2.5, 4.0])

    doc.add_heading("2.3 LLM Quality Assessment", level=2)
    doc.add_paragraph(
        "Each passage is evaluated by a large language model in a single assessment call. "
        "Three quality dimensions are scored:"
    )
    llm_rows = [
        ("Sensitivity", "0.30 maximum", "Excluded if score exceeds threshold", "Detects hate speech, explicit violence, sexual content, and similar unsafe material (scale: 0 = safe, 1 = highly sensitive)"),
        ("Coherence", "0.60 minimum", "Excluded if score falls below threshold", "Measures whether text reads as natural, coherent prose versus incoherent scraped content (scale: 0 = incoherent, 1 = fully coherent)"),
        ("Domain classification", "—", "Assigned, not used for exclusion", "Categorizes content into a standardized domain taxonomy for dataset organization"),
        ("Prohibited content", "—", "Always excluded", "Content flagged as prohibited by the LLM safety system is removed"),
    ]
    add_table_from_rows(
        doc,
        ["Assessment", "Threshold", "Action", "Description"],
        llm_rows,
        col_widths=[1.2, 1.0, 1.3, 3.0],
    )

    doc.add_paragraph("Content domain taxonomy:")
    domains = [
        "Education", "Public domain reference", "News and journalism", "Stories and literature",
        "General information", "Biography", "Science and technology", "Legal and government",
        "Religion and culture", "Other (custom label when none apply)",
    ]
    for d in domains:
        doc.add_paragraph(d, style="List Bullet")

    doc.add_heading("2.4 Paragraph Length Standardization", level=2)
    doc.add_paragraph(
        "After quality filtering, passages are standardized to a target length suitable "
        "for TTS recording sessions. Related fragments from the same source document "
        "are consolidated before length assessment."
    )
    length_rows = [
        ("Minimum word count", "300 words", "Passages below this threshold are excluded"),
        ("Maximum word count", "450 words", "Hard upper limit per paragraph chunk"),
        ("Pruning range", "451–1,000 words", "Trimmed to 450 words at natural sentence boundaries"),
        ("Splitting threshold", "Above 1,000 words", "Split into multiple 300–450 word chunks at sentence boundaries"),
        ("Word counting method", "Whitespace-based", "Consistent with curation standards; Indic-aware tokenization available as alternative"),
        ("Sentence boundary detection", "Cross-script", "Recognizes Indic danda, Arabic full stop, and standard Latin punctuation"),
    ]
    add_table_from_rows(doc, ["Parameter", "Value", "Action"], length_rows, col_widths=[2.0, 1.5, 3.2])


def add_phase2(doc: Document) -> None:
    doc.add_heading("3. Phase 2 — TTS Readiness Optimization", level=1)
    doc.add_paragraph(
        "Phase 2 prepares validated text for speech synthesis by removing non-standard "
        "words (NSW) — numerical expressions, citations, measurements, and structural "
        "artifacts that TTS engines cannot render naturally in Indic languages. "
        "The target output remains 300–450 words per paragraph."
    )

    doc.add_heading("3.1 Structural Content Filtering", level=2)
    doc.add_paragraph(
        "Passages with heavy list, table, or non-prose structure are excluded before "
        "NSW processing, as they are unsuitable for natural speech recording."
    )
    struct_rows = [
        ("List-heavy content", "Excluded when ≥ 5 list/table markers are present"),
        ("Table-like content", "Excluded when ≥ 6 pipe (|) column separators are present"),
        ("Numbered list dominance", "Excluded when ≥ 40% of lines are numbered list items (minimum 4 lines)"),
        ("Empty after cleanup", "Excluded when all content is removed during processing"),
    ]
    add_table_from_rows(doc, ["Filter", "Criterion"], struct_rows, col_widths=[2.5, 4.0])

    doc.add_heading("3.2 Text Preprocessing", level=2)
    doc.add_paragraph("Before NSW removal, each passage undergoes the following preprocessing:")
    pre_rows = [
        ("1", "Encoding repair", "Corrects text encoding artifacts"),
        ("2", "Whitespace normalization", "Standardizes spaces; removes zero-width and non-breaking characters"),
        ("3", "Placeholder removal", "Strips box and checkbox placeholder characters"),
        ("4", "Wiki markup removal", "Removes encyclopedic link and template markup"),
        ("5", "Parenthetical content removal", "Strips parenthetical blocks that typically contain citations or references"),
        ("6", "Boilerplate block removal", "Removes reference sections, bibliographies, footnotes, and horizontal rule lines"),
        ("7", "Cosmetic cleanup", "Collapses excess whitespace; fixes orphan punctuation and dangling connectors"),
    ]
    add_table_from_rows(doc, ["#", "Step", "Description"], pre_rows, col_widths=[0.4, 1.8, 4.3])

    doc.add_heading("3.3 Non-Standard Word (NSW) Removal", level=2)
    doc.add_paragraph(
        "A comprehensive pattern library removes TTS-hostile content. Matches are "
        "replaced with whitespace; passages that become empty are excluded. "
        "Patterns are organized into four tiers:"
    )

    doc.add_heading("Tier 1 — Compound and Technical Patterns", level=3)
    tier1 = [
        "ISBN, ISSN, and DOI identifiers",
        "IP addresses and hex/color codes",
        "File paths and document extensions",
        "HTML entities",
        "Geographic coordinates and degree-minute-second notation",
        "Temperature values",
        "Scientific and chemical notation",
        "Subscript and superscript characters",
        "Bank and card number patterns",
        "Currency amounts (₹, Rs., $, €, and written forms)",
        "Large comma-separated numbers",
    ]
    for item in tier1:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Tier 2 — Measurements, Time, and References", level=3)
    tier2 = [
        "Measurement units with numeric values (distance, weight, pressure, frequency, etc.)",
        "Percentages",
        "Time zones and UTC/GMT offsets",
        "Clock times and durations",
        "Dates and year ranges",
        "Ratios and scale notations",
        "Version and model numbers",
        "Page, chapter, and lesson references (including Indic script labels)",
        "Figure, table, and image references (including Indic script labels)",
    ]
    for item in tier2:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Tier 3 — Citations, Symbols, and Ordinals", level=3)
    tier3 = [
        "Roman numerals",
        "Citation brackets and footnote markers",
        "Ordinal numbers (English and Indic forms)",
        "PIN and postal codes",
        "Vehicle registration numbers",
        "Author-year citations and bibliographic references",
        "Hashtags and social media mentions",
        "Copyright and trademark symbols",
        "Mathematical operators and expressions",
        "ASCII emoticons and kaomoji",
    ]
    for item in tier3:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Tier 4 — Residual Digit and Symbol Cleanup", level=3)
    tier4 = [
        "Decimal numbers",
        "Degree symbols",
        "Inline bullet characters",
        "Backtick, pipe, and asterisk runs",
        "Remaining standalone integers",
    ]
    for item in tier4:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("3.4 Final Paragraph Sizing", level=2)
    doc.add_paragraph(
        "After NSW removal, passages are re-sized to meet TTS recording requirements:"
    )
    sizing_rows = [
        ("Up to 450 words", "Retained as a single paragraph"),
        ("451–599 words", "Split at a sentence boundary (target 300–400 words); remainder routed separately"),
        ("600 words and above", "Recursively bisected at sentence boundaries until all chunks meet size limits"),
    ]
    add_table_from_rows(doc, ["Passage length", "Treatment"], sizing_rows, col_widths=[2.0, 4.5])

    doc.add_paragraph(
        "Final dataset passages are those within the 300–450 word range after all "
        "Phase 2 processing is complete."
    )


def add_comparison(doc: Document) -> None:
    doc.add_heading("4. Phase Comparison Summary", level=1)
    rows = [
        ("Primary objective", "Content quality and safety assurance", "TTS synthesis readiness"),
        ("LLM assessment", "Yes — coherence, sensitivity, domain", "No — rule-based processing only"),
        ("Target word range", "300–450 words", "300–450 words (re-enforced)"),
        ("PII handling", "Detected and redacted", "Removed via digit pattern matching"),
        ("URL and link handling", "Removed from text", "File paths and wiki URLs removed"),
        ("List and table content", "Retained if coherent", "Excluded when list-heavy"),
        ("Numerical content", "Retained", "Comprehensively removed"),
        ("Citations and references", "Retained", "Removed"),
    ]
    add_table_from_rows(doc, ["Aspect", "Phase 1", "Phase 2"], rows, col_widths=[2.0, 2.5, 2.5])


def add_conclusion(doc: Document) -> None:
    doc.add_heading("5. Conclusion", level=1)
    doc.add_paragraph(
        "The two-phase validation framework provides systematic quality assurance for "
        "Indic-language TTS datasets. Phase 1 ensures that retained content is "
        "linguistically coherent, culturally appropriate, privacy-compliant, and "
        "structurally sound. Phase 2 transforms validated prose into speech-ready "
        "passages by eliminating numerical expressions, citations, and non-prose "
        "artifacts that would impede natural TTS rendering."
    )
    doc.add_paragraph(
        "Together, these phases produce standardized 300–450 word paragraphs suitable "
        "for professional voice recording across 15 Indic languages in the curated corpus."
    )


def main() -> None:
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    add_title_page(doc)
    add_executive_summary(doc)
    add_section_overview(doc)
    add_phase1(doc)
    add_phase2(doc)
    add_comparison(doc)
    add_conclusion(doc)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
