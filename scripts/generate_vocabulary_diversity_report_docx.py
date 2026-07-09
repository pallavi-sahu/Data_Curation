#!/usr/bin/env python3
"""Generate client-facing unique-word diversity report DOCX from JSON summary."""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

_REPO_ROOT = Path(__file__).resolve().parent.parent
JSON_PATH = _REPO_ROOT / "final" / "vocabulary_diversity_reports.json"
OUTPUT_PATH = _REPO_ROOT / "final" / "vocabulary_diversity_report.docx"


def _set_cell_shading(cell, fill: str) -> None:
    shading = cell._element.get_or_add_tcPr()
    shd = shading.makeelement(
        qn("w:shd"),
        {qn("w:fill"): fill, qn("w:val"): "clear"},
    )
    shading.append(shd)


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    *,
    col_widths: list[float] | None = None,
) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    for i, text in enumerate(headers):
        hdr.cells[i].text = text
        _set_cell_shading(hdr.cells[i], "D6E4F0")
        for p in hdr.cells[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(10)

    for r_idx, row_data in enumerate(rows):
        row = table.rows[r_idx + 1]
        for c_idx, text in enumerate(row_data):
            row.cells[c_idx].text = text
            for p in row.cells[c_idx].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)

    if col_widths:
        for row in table.rows:
            for i, width in enumerate(col_widths):
                row.cells[i].width = Inches(width)

    doc.add_paragraph()


def _display_name(lang: str) -> str:
    return "Assamese" if lang == "Assamese" else lang.capitalize()


def _mean_unique_pct(report: dict) -> float:
    value = report["unique_word_percentage_per_script"]["mean"]
    return value if value is not None else 0.0


def add_title_page(doc: Document, *, languages: int, total_scripts: int) -> None:
    for _ in range(4):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Unique Word Diversity Report")
    run.bold = True
    run.font.size = Pt(28)
    run.font.name = "Calibri"
    run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Percentage of Unique Words Across Indic Language Corpora")
    run.font.size = Pt(14)
    run.font.name = "Calibri"
    run.font.color.rgb = RGBColor(0x44, 0x54, 0x6A)

    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(
        f"Report Date: {date.today():%d %B %Y}  |  "
        f"{languages} languages  |  {total_scripts:,} scripts"
    )
    run.font.size = Pt(11)
    run.font.name = "Calibri"
    run.italic = True

    doc.add_page_break()


def add_executive_summary(doc: Document, reports: dict) -> None:
    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(
        "This report measures unique-word diversity in the live curated TTS corpora. "
        "For each language it reports the average percentage of unique words per script, "
        "along with corpus-wide unique-word counts and the share of scripts below the "
        "25% uniqueness threshold used in validation."
    )

    total_scripts = sum(r["scripts"] for r in reports.values())
    total_words = sum(r["total_words"] for r in reports.values())
    ranked = sorted(reports.items(), key=lambda x: _mean_unique_pct(x[1]), reverse=True)
    highest = ranked[0]
    lowest = ranked[-1]

    add_table(
        doc,
        ["Metric", "Value"],
        [
            ("Languages covered", str(len(reports))),
            ("Total scripts", f"{total_scripts:,}"),
            ("Total words (all languages)", f"{total_words:,}"),
            (
                "Highest avg unique word % (per script)",
                f"{_display_name(highest[0])} — {_mean_unique_pct(highest[1]):.1f}%",
            ),
            (
                "Lowest avg unique word % (per script)",
                f"{_display_name(lowest[0])} — {_mean_unique_pct(lowest[1]):.1f}%",
            ),
        ],
        col_widths=[3.2, 3.3],
    )


def add_summary_table(doc: Document, reports: dict) -> None:
    doc.add_heading("1. Unique Word % by Language", level=1)
    doc.add_paragraph(
        "Languages are sorted by average unique word percentage per script "
        "(unique words in a script ÷ total words in that script)."
    )

    rows = []
    for lang in sorted(reports, key=lambda name: _mean_unique_pct(reports[name]), reverse=True):
        r = reports[lang]
        per_script = r["unique_word_percentage_per_script"]
        rows.append(
            [
                _display_name(lang),
                f"{r['scripts']:,}",
                f"{r['total_words']:,}",
                f"{r['unique_words']:,}",
                f"{per_script['mean']:.1f}%",
                f"{r['unique_word_percentage_corpus']:.2f}%",
                f"{r['scripts_below_25_percent']:,} ({r['scripts_below_25_percent_pct']:.1f}%)",
            ]
        )

    add_table(
        doc,
        [
            "Language",
            "Scripts",
            "Total words",
            "Unique words",
            "Avg unique word %",
            "Corpus unique word %",
            "Scripts <25%",
        ],
        rows,
        col_widths=[0.9, 0.55, 0.75, 0.75, 0.85, 0.85, 0.75],
    )


def add_language_detail(doc: Document, lang: str, report: dict) -> None:
    per_script = report["unique_word_percentage_per_script"]
    doc.add_heading(f"{_display_name(lang)}", level=2)
    doc.add_paragraph(
        f"{report['scripts']:,} scripts · {report['total_words']:,} total words · "
        f"{report['unique_words']:,} unique words"
    )

    add_table(
        doc,
        ["Metric", "Value"],
        [
            ("Avg unique word % (per script)", f"{per_script['mean']:.1f}%"),
            ("Median unique word % (per script)", f"{per_script['median']:.1f}%"),
            (
                "Min / max unique word % (per script)",
                f"{per_script['min']:.1f}% / {per_script['max']:.1f}%",
            ),
            ("Corpus unique word %", f"{report['unique_word_percentage_corpus']:.2f}%"),
            ("Mean words per script", str(report["mean_words_per_script"])),
            (
                "Scripts below 25% unique words",
                f"{report['scripts_below_25_percent']:,} ({report['scripts_below_25_percent_pct']:.1f}%)",
            ),
        ],
        col_widths=[2.8, 3.7],
    )


def add_methodology(doc: Document) -> None:
    doc.add_heading("2. Methodology", level=1)
    doc.add_paragraph(
        "Tokenization follows the validation pipeline (Indic-aware word boundaries). "
        "For each script, unique word % = (number of distinct words in the script ÷ total "
        "words in the script) × 100. The reported average is the mean of this value across "
        "all scripts in a language.",
        style="List Bullet",
    )
    doc.add_paragraph(
        "Corpus unique word % counts each distinct word once across the entire language "
        "corpus, divided by the total word count across all scripts.",
        style="List Bullet",
    )
    doc.add_paragraph(
        "Full metrics are available in final/vocabulary_diversity_reports.json.",
    )


def main() -> None:
    payload = json.loads(JSON_PATH.read_text(encoding="utf-8"))
    reports: dict = payload["languages"]
    total_scripts = sum(r["scripts"] for r in reports.values())

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    add_title_page(doc, languages=len(reports), total_scripts=total_scripts)
    add_executive_summary(doc, reports)
    add_summary_table(doc, reports)

    doc.add_page_break()
    doc.add_heading("Appendix — Per-Language Detail", level=1)
    for lang in sorted(reports, key=lambda name: _mean_unique_pct(reports[name]), reverse=True):
        add_language_detail(doc, lang, reports[lang])

    doc.add_page_break()
    add_methodology(doc)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
